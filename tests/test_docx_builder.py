from docx import Document
from src import docx_builder

def test_build_heading_meta_paragraphs(tmp_path):
    p = str(tmp_path / "001_test.docx")
    meta = {"Publisher": "Kompas.com", "Date": "2025-03-10", "URL": "https://x.com/1", "Keyword_Found": "k1", "Accessed": "2026-09-24"}
    body = "Paragraf satu tentang Microsoft.\n\nParagraf dua tentang Indonesia."
    docx_builder.build(p, "Judul Artikel", meta, body)
    doc = Document(p)
    assert doc.paragraphs[0].style.name.startswith("Heading 1")
    assert "Judul Artikel" in doc.paragraphs[0].text
    texts = [x.text for x in doc.paragraphs]
    assert any("Kompas.com" in t and "2025-03-10" in t for t in texts)
    assert any("Paragraf satu" in t for t in texts)
    assert any("Paragraf dua" in t for t in texts)

def test_build_rejects_empty_body(tmp_path):
    import pytest
    with pytest.raises(ValueError):
        docx_builder.build(str(tmp_path / "x.docx"), "T", {}, "   ")
