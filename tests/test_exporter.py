from src import exporter

def test_export_creates_xlsx_with_headers(tmp_path):
    rows = [{"Publisher": "Kompas", "Title_ID": "Judul ID", "Title_Original": "Title EN", "url": "https://x.com/1", "URL": "https://x.com/1", "date_iso": "2025-03-10", "Date": "2025-03-10", "Keyword_Found": "k1", "Also_Found_In": "", "Snippet_ID": "S ID", "Snippet_Original": "S EN", "Source_Lang": "en", "translation_ok": True}]
    p = str(tmp_path / "out.xlsx")
    exporter.export_xlsx(rows, p)
    import openpyxl
    wb = openpyxl.load_workbook(p)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert headers[0] == "No"
    assert "Publisher" in headers
    assert "Title_ID" in headers
    assert ws.max_row == 2
