"""XLSX + run-meta writer."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

COLUMNS = ["No", "Publisher", "Title_ID", "Title_Original", "URL", "Date", "Keyword_Found", "Also_Found_In", "Snippet_ID", "Snippet_Original", "Source_Lang", "translation_ok"]
WIDTHS = {"No": 6, "Publisher": 18, "Title_ID": 60, "Title_Original": 60, "URL": 50, "Date": 12, "Keyword_Found": 30, "Also_Found_In": 30, "Snippet_ID": 60, "Snippet_Original": 60, "Source_Lang": 12, "translation_ok": 14}

def export_xlsx(rows, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "articles-2025"
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(bold=True, color="FFFFFF")
    for j, col in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=j, value=col)
        c.fill = hdr_fill
        c.font = hdr_font
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, r in enumerate(rows, start=1):
        ws.cell(row=i + 1, column=1, value=i)
        ws.cell(row=i + 1, column=2, value=r.get("Publisher") or r.get("publisher", ""))
        ws.cell(row=i + 1, column=3, value=r.get("Title_ID", ""))
        ws.cell(row=i + 1, column=4, value=r.get("Title_Original", ""))
        ws.cell(row=i + 1, column=5, value=r.get("URL") or r.get("url", ""))
        ws.cell(row=i + 1, column=6, value=r.get("Date") or r.get("date_iso", ""))
        ws.cell(row=i + 1, column=7, value=r.get("Keyword_Found", ""))
        ws.cell(row=i + 1, column=8, value=r.get("Also_Found_In", ""))
        ws.cell(row=i + 1, column=9, value=r.get("Snippet_ID", ""))
        ws.cell(row=i + 1, column=10, value=r.get("Snippet_Original", ""))
        ws.cell(row=i + 1, column=11, value=r.get("Source_Lang", ""))
        ws.cell(row=i + 1, column=12, value=str(r.get("translation_ok", "")))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for j, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = WIDTHS[col]
    wb.save(path)
    return path

def write_meta(meta, path):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(meta, f, ensure_ascii=False, indent=2)
    return path
