"""Assemble one Indonesian docx: Heading-1 title, metadata, body paragraphs."""
from docx import Document

def build(path, title, meta, body_id):
    body = (body_id or "").strip()
    if not body:
        raise ValueError("empty Indonesian body")
    doc = Document()
    doc.add_heading((title or "Tanpa judul").strip(), level=1)
    doc.add_paragraph(
        f"Publisher: {meta.get('Publisher', '')} | "
        f"Date: {meta.get('Date', '')} | "
        f"Keyword: {meta.get('Keyword_Found', '')} | "
        f"Accessed: {meta.get('Accessed', '')}"
    )
    doc.add_paragraph(f"URL: {meta.get('URL', '')}")
    for para in body.split("\n"):
        para = para.strip()
        if para:
            doc.add_paragraph(para)
    doc.save(path)
    return path
